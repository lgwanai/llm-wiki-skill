#!/usr/bin/env python3
"""ledger.py — Structured table management backed by DuckDB.

Stores tables in .wiki/ledger/ledger.duckdb:
  - _registry meta-table: display_name → actual_name mapping + schema info
  - One DuckDB table per user table, with typed columns
  - SEQUENCE per table for auto-increment
  - _embeddings table for vector search over table rows

Usage (via wiki.py):
    wiki ledger list | show | create | insert | update-schema | delete | stats | embed

Auto-migrates from old JSON format on first use if JSON files exist.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

import duckdb

# Ensure scripts/ is on sys.path for config import
sys.path.insert(0, str(Path(__file__).resolve().parent))

from config import get_wiki_dir  # noqa: E402

WIKI_DIR = get_wiki_dir()
LEDGER_DIR = WIKI_DIR / "ledger"
LEDGER_DB = LEDGER_DIR / "ledger.duckdb"
# Legacy JSON paths (kept for auto-migration)
TABLES_DIR = LEDGER_DIR / "tables"
REGISTRY_FILE = LEDGER_DIR / "registry.json"
INDEX_FILE = LEDGER_DIR / "index.json"

VALID_TYPES = frozenset({"string", "text", "integer", "number", "boolean", "date", "datetime"})

TYPE_MAP: dict[str, str] = {
    "string": "VARCHAR",
    "text": "VARCHAR",
    "integer": "INTEGER",
    "number": "DOUBLE",
    "boolean": "BOOLEAN",
    "date": "DATE",
    "datetime": "TIMESTAMP",
}

TYPE_HINTS = {
    "integer": r"^\s*-?\d+\s*$",
    "number": r"^\s*-?\d+\.?\d*\s*$",
    "percentage": r"^\s*-?\d+\.?\d*\s*%\s*$",
    "date": r"^\s*\d{4}[-/]\d{1,2}[-/]\d{1,2}\s*$",
    "boolean": r"^\s*(true|false|yes|no|是|否|0|1)\s*$",
}


# ═══════════════════════════════════════════════════════════════════════
# Helpers
# ═══════════════════════════════════════════════════════════════════════

def _slugify(name: str) -> str:
    s = name.lower().strip().replace(" ", "-").replace("_", "-")
    s = re.sub(r"[^a-z0-9-]", "", s)
    s = re.sub(r"-+", "-", s)
    return s.strip("-")


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _q(name: str) -> str:
    """Safely quote a SQL identifier by escaping embedded double-quotes."""
    return '"' + name.replace('"', '""') + '"'


def _get_conn(read_only: bool = False) -> duckdb.DuckDBPyConnection:
    """Get a DuckDB connection, ensuring schema tables and auto-migration."""
    LEDGER_DB.parent.mkdir(parents=True, exist_ok=True)
    if read_only:
        if not LEDGER_DB.exists():
            raise FileNotFoundError("No ledger database found.")
        return duckdb.connect(str(LEDGER_DB), read_only=True)
    conn = duckdb.connect(str(LEDGER_DB))
    _ensure_schema(conn)
    # Auto-migrate from JSON if needed
    _auto_migrate(conn)
    return conn


def _ensure_schema(conn: duckdb.DuckDBPyConnection) -> None:
    """Create _registry and _embeddings tables if they don't exist."""
    conn.execute("""
        CREATE TABLE IF NOT EXISTS _registry (
            actual_name        VARCHAR PRIMARY KEY,
            display_name       VARCHAR NOT NULL,
            description        VARCHAR DEFAULT '',
            record_count       INTEGER DEFAULT 0,
            created_at         TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            updated_at         TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            fields_json        VARCHAR DEFAULT '[]',
            unique_key         VARCHAR DEFAULT '[]',
            auto_increment     BOOLEAN DEFAULT FALSE,
            auto_increment_field VARCHAR DEFAULT NULL
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS _embeddings (
            table_name  VARCHAR NOT NULL,
            row_id      INTEGER NOT NULL,
            embedding   FLOAT[],   -- dimension depends on model
            updated_at  TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            PRIMARY KEY (table_name, row_id)
        )
    """)


def _auto_migrate(conn: duckdb.DuckDBPyConnection) -> None:
    """One-time migration from old JSON ledger to DuckDB."""
    existing = conn.execute("SELECT COUNT(*) FROM _registry").fetchone()[0]
    if existing == 0 and REGISTRY_FILE.exists():
        # Old v1 registry layout: .wiki/ledger/tables/<table>/*
        _migrate_from_json(conn)
    if INDEX_FILE.exists():
        # Later import layout: .wiki/ledger/<table-id>/{schema.json,data.json}
        _migrate_from_index_json(conn)


def _migrate_from_json(conn: duckdb.DuckDBPyConnection) -> None:
    """Migrate all JSON-based tables into DuckDB."""
    reg = _load_json_legacy(REGISTRY_FILE, {"version": 1, "tables": {}})
    tables = reg.get("tables", {})
    if not tables:
        return

    for actual_name, info in tables.items():
        schema = _load_json_legacy(TABLES_DIR / actual_name / "schema.json")
        data = _load_json_legacy(TABLES_DIR / actual_name / "data.json", [])
        seq = _load_json_legacy(TABLES_DIR / actual_name / "sequence.json", {"next_id": 1})
        if not schema:
            continue

        fields = schema.get("fields", [])
        # Build column DDL
        col_defs = []
        for fdef in fields:
            duck_type = TYPE_MAP.get(fdef.get("type", "string"), "VARCHAR")
            col_defs.append(f'"{fdef["name"]}" {duck_type}')
        conn.execute(f'CREATE TABLE IF NOT EXISTS {_q(actual_name)} ({", ".join(col_defs)})')

        # Sequence
        if schema.get("auto_increment"):
            next_id = seq.get("next_id", 1)
            if data:
                max_id = max((row.get("_id", 0) for row in data), default=0)
                next_id = max(next_id, max_id + 1)
            conn.execute(f'CREATE SEQUENCE IF NOT EXISTS {_q(f"seq_{actual_name}")} START {next_id}')

        # Insert data
        if data:
            col_names = [f["name"] for f in fields]
            quoted_cols = [f'"{c}"' for c in col_names]
            placeholders = ", ".join(["?" for _ in col_names])
            for row in data:
                vals = [row.get(c) for c in col_names]
                try:
                    conn.execute(
                        f'INSERT INTO {_q(actual_name)} ({", ".join(quoted_cols)}) VALUES ({placeholders})',
                        vals,
                    )
                except duckdb.Error:
                    pass  # Skip rows that fail migration

        # Register
        conn.execute(
            """INSERT INTO _registry
               (actual_name, display_name, description, record_count, created_at, updated_at,
                fields_json, unique_key, auto_increment, auto_increment_field)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            [
                actual_name,
                schema.get("display_name", info.get("display_name", actual_name)),
                schema.get("description", info.get("description", "")),
                len(data) if data else 0,
                schema.get("created_at", info.get("created_at", _now_iso())),
                schema.get("updated_at", info.get("updated_at", _now_iso())),
                json.dumps(schema.get("fields", []), ensure_ascii=False),
                json.dumps(schema.get("unique_key", [])),
                schema.get("auto_increment", False),
                schema.get("auto_increment_field"),
            ],
        )


def _migrate_from_index_json(conn: duckdb.DuckDBPyConnection) -> None:
    """Migrate index.json based imported ledgers into DuckDB."""
    entries = _load_json_legacy(INDEX_FILE, [])
    if not isinstance(entries, list):
        return

    existing = {
        r[0] for r in conn.execute("SELECT actual_name FROM _registry").fetchall()
    }

    for entry in entries:
        if not isinstance(entry, dict):
            continue
        table_id = entry.get("id")
        if not table_id or table_id in existing:
            continue

        table_dir = LEDGER_DIR / table_id
        schema = _load_json_legacy(table_dir / "schema.json")
        data = _load_json_legacy(table_dir / "data.json", [])
        if not schema or not isinstance(data, list):
            continue

        fields = []
        for fdef in schema.get("fields", []):
            display_name = fdef.get("display_name") or fdef.get("name")
            if not display_name:
                continue
            fields.append({
                "name": display_name,
                "type": _normalize_import_type(fdef.get("type", "string")),
                "required": False,
                "description": "",
                "source_name": fdef.get("name", display_name),
            })
        if not fields:
            continue

        actual_name = table_id if re.match(r"^[A-Za-z_][A-Za-z0-9_-]*$", table_id) else _generate_table_name(
            schema.get("name", table_id), conn
        )
        col_defs = [f'{_q(f["name"])} {TYPE_MAP.get(f["type"], "VARCHAR")}' for f in fields]
        try:
            conn.execute(f'CREATE TABLE IF NOT EXISTS {_q(actual_name)} ({", ".join(col_defs)})')
        except duckdb.Error:
            continue

        col_names = [f["name"] for f in fields]
        source_names = [f.get("source_name", f["name"]) for f in fields]
        quoted_cols = [_q(c) for c in col_names]
        placeholders = ", ".join(["?" for _ in col_names])
        inserted = 0
        for row in data:
            if not isinstance(row, dict):
                continue
            vals = []
            for field, source_name in zip(fields, source_names):
                raw = row.get(source_name, row.get(field["name"]))
                vals.append(_clean_import_value(raw, field["type"]))
            try:
                conn.execute(
                    f'INSERT INTO {_q(actual_name)} ({", ".join(quoted_cols)}) VALUES ({placeholders})',
                    vals,
                )
                inserted += 1
            except duckdb.Error:
                pass

        conn.execute(
            """INSERT INTO _registry
               (actual_name, display_name, description, record_count, created_at, updated_at,
                fields_json, unique_key, auto_increment, auto_increment_field)
               VALUES (?, ?, ?, ?, ?, ?, ?, '[]', FALSE, NULL)""",
            [
                actual_name,
                schema.get("name", entry.get("name", actual_name)),
                schema.get("description", entry.get("description", "")),
                inserted,
                schema.get("import_time", _now_iso()),
                schema.get("import_time", _now_iso()),
                json.dumps([{k: v for k, v in f.items() if k != "source_name"} for f in fields], ensure_ascii=False),
            ],
        )
        existing.add(actual_name)


def _load_json_legacy(path: Path, default=None):
    """Load a JSON file, returning *default* if missing."""
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except (FileNotFoundError, json.JSONDecodeError, OSError):
        return default


# ═══════════════════════════════════════════════════════════════════════
# Table name generation
# ═══════════════════════════════════════════════════════════════════════

def _generate_table_name(display_name: str, conn: duckdb.DuckDBPyConnection) -> str:
    """Generate a unique, English-safe table name."""
    base = _slugify(display_name)
    if not base:
        h = hashlib.sha256(display_name.encode("utf-8")).hexdigest()[:8]
        base = f"table_{h}"

    existing = {
        r[0] for r in conn.execute("SELECT actual_name FROM _registry").fetchall()
    }
    candidate = base
    suffix = 1
    while candidate in existing:
        suffix += 1
        candidate = f"{base}_{suffix}"
    return candidate


def _resolve_table(user_input: str, conn: duckdb.DuckDBPyConnection) -> str | None:
    """Look up *user_input* in _registry. Returns actual_name or None."""
    # Exact actual_name
    row = conn.execute("SELECT actual_name FROM _registry WHERE actual_name = ?", [user_input]).fetchone()
    if row:
        return row[0]
    # Exact display_name
    row = conn.execute("SELECT actual_name FROM _registry WHERE display_name = ?", [user_input]).fetchone()
    if row:
        return row[0]
    # Case-insensitive display_name
    row = conn.execute(
        "SELECT actual_name FROM _registry WHERE LOWER(display_name) = LOWER(?)", [user_input]
    ).fetchone()
    return row[0] if row else None


# ═══════════════════════════════════════════════════════════════════════
# Type validation & coercion (unchanged from JSON version)
# ═══════════════════════════════════════════════════════════════════════

def _coerce_value(value, field_type: str):
    if value is None:
        return None, None
    if field_type in ("string", "text"):
        return str(value), None
    if field_type == "integer":
        if isinstance(value, bool):
            return value, f"Expected integer, got boolean"
        try:
            return int(value), None
        except (ValueError, TypeError):
            return value, f"Expected integer, got '{value}'"
    if field_type == "number":
        if isinstance(value, bool):
            return value, f"Expected number, got boolean"
        try:
            return float(value), None
        except (ValueError, TypeError):
            return value, f"Expected number, got '{value}'"
    if field_type == "boolean":
        if isinstance(value, bool):
            return value, None
        if isinstance(value, str):
            low = value.strip().lower()
            if low in ("true", "1", "yes", "y", "是"):
                return True, None
            if low in ("false", "0", "no", "n", "", "否"):
                return False, None
        if isinstance(value, (int, float)):
            return bool(value), None
        return value, f"Cannot coerce '{value}' to boolean"
    if field_type == "date":
        s = str(value).strip()
        if re.match(r"^\d{4}-\d{2}-\d{2}$", s):
            return s, None
        return value, f"Expected date (YYYY-MM-DD), got '{value}'"
    if field_type == "datetime":
        s = str(value).strip()
        if re.match(r"^\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}", s):
            return s, None
        return value, f"Expected datetime (ISO 8601), got '{value}'"
    return value, f"Unknown type '{field_type}'"


def _normalize_import_type(type_name: str) -> str:
    """Map inferred/imported types to ledger field types."""
    if type_name in VALID_TYPES:
        return type_name
    if type_name in ("int", "integer"):
        return "integer"
    if type_name in ("float", "number", "percentage"):
        return "number"
    if type_name in ("url", "email", "number_range"):
        return "string"
    if type_name == "bool":
        return "boolean"
    return "text"


def _clean_import_value(value, field_type: str):
    """Normalize imported cell values before inserting into DuckDB."""
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        if stripped == "":
            return None
        value = stripped
    if field_type == "number" and isinstance(value, str):
        value = value.replace(",", "").replace("%", "")
    if field_type == "integer" and isinstance(value, str):
        value = value.replace(",", "")
    if field_type == "date" and isinstance(value, str):
        value = value.replace("/", "-")
    coerced, err = _coerce_value(value, field_type)
    return value if err else coerced


def _infer_import_type(values: list[str]) -> str:
    non_empty = [str(v).strip() for v in values if str(v).strip()]
    if not non_empty:
        return "text"
    for type_name, pattern in TYPE_HINTS.items():
        if all(re.match(pattern, v) for v in non_empty):
            return _normalize_import_type(type_name)
    int_count = sum(1 for v in non_empty if re.match(TYPE_HINTS["integer"], v))
    if int_count / len(non_empty) > 0.8:
        return "integer"
    number_count = sum(1 for v in non_empty if re.match(TYPE_HINTS["number"], v))
    if number_count / len(non_empty) > 0.8:
        return "number"
    return "text"


def _dedupe_headers(headers: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    result = []
    for idx, raw in enumerate(headers, start=1):
        base = str(raw).strip() or f"column_{idx}"
        count = seen.get(base, 0) + 1
        seen[base] = count
        result.append(base if count == 1 else f"{base}_{count}")
    return result


def _validate_row(row: dict, schema: dict, conn, actual_name: str) -> list[dict]:
    """Validate a single row against *schema*. Checks unique constraints via SQL."""
    errors: list[dict] = []
    fields = {f["name"]: f for f in schema.get("fields", [])}
    unique_key = schema.get("unique_key", [])

    # Unknown fields
    for key in row:
        if key not in fields:
            errors.append({"field": key, "error": f"Unknown field '{key}'"})

    coerced: dict = {}
    for fdef in schema.get("fields", []):
        fname = fdef["name"]
        ftype = fdef.get("type", "string")
        required = fdef.get("required", False)
        auto_inc = fdef.get("auto_increment", False)
        if auto_inc:
            continue

        raw = row.get(fname)
        if required and (raw is None or (isinstance(raw, str) and raw.strip() == "")):
            errors.append({"field": fname, "error": f"'{fname}' is required"})
            continue

        if raw is not None:
            val, err = _coerce_value(raw, ftype)
            if err:
                errors.append({"field": fname, "error": err})
            else:
                coerced[fname] = val
        else:
            coerced[fname] = None

    if errors:
        return errors

    # Unique constraint check via SQL
    for uk_field in unique_key:
        new_val = coerced.get(uk_field)
        if new_val is None:
            continue
        try:
            row_exists = conn.execute(
                f'SELECT 1 FROM {_q(actual_name)} WHERE {_q(uk_field)} = ? LIMIT 1', [new_val]
            ).fetchone()
            if row_exists:
                errors.append({
                    "field": uk_field,
                    "error": f"Duplicate value '{new_val}' (already exists in table)",
                })
        except duckdb.Error:
            pass

    return errors


# ═══════════════════════════════════════════════════════════════════════
# Command handlers
# ═══════════════════════════════════════════════════════════════════════

def cmd_list() -> dict:
    conn = _get_conn()
    rows = conn.execute(
        "SELECT actual_name, display_name, description, record_count, created_at, updated_at "
        "FROM _registry ORDER BY display_name"
    ).fetchall()
    tables = []
    for r in rows:
        tables.append({
            "display_name": r[1],
            "actual_name": r[0],
            "description": r[2] or "",
            "record_count": r[3],
            "created_at": r[4].isoformat() if r[4] else "",
            "updated_at": r[5].isoformat() if r[5] else "",
        })
    return {"success": True, "count": len(tables), "tables": tables}


def cmd_show(table: str) -> dict:
    try:
        conn = _get_conn(read_only=True)
    except FileNotFoundError:
        return {"success": False, "error": f"Table '{table}' not found. Use 'wiki ledger list'."}
    actual = _resolve_table(table, conn)
    if actual is None:
        return {"success": False, "error": f"Table '{table}' not found. Use 'wiki ledger list'."}

    reg = conn.execute(
        "SELECT display_name, description, fields_json, unique_key, auto_increment, "
        "auto_increment_field, created_at, updated_at FROM _registry WHERE actual_name = ?",
        [actual],
    ).fetchone()
    if reg is None:
        return {"success": False, "error": f"Registry entry missing for '{actual}'."}

    fields = json.loads(reg[2])
    unique_key = json.loads(reg[3])

    # Get rows
    try:
        rows = conn.execute(f'SELECT * FROM {_q(actual)} LIMIT 20').fetchall()
        col_names = [desc[0] for desc in conn.description]
        data = [dict(zip(col_names, r)) for r in rows]
        total = conn.execute(f'SELECT COUNT(*) FROM {_q(actual)}').fetchone()[0]
    except duckdb.Error as e:
        return {"success": False, "error": f"Failed to read table: {e}"}

    return {
        "success": True,
        "table": {
            "display_name": reg[0],
            "actual_name": actual,
            "description": reg[1] or "",
            "auto_increment": reg[4],
            "auto_increment_field": reg[5],
            "unique_key": unique_key,
            "fields": fields,
            "created_at": reg[6].isoformat() if reg[6] else "",
            "updated_at": reg[7].isoformat() if reg[7] else "",
        },
        "data": data,
        "total_rows": total,
        "shown_rows": len(data),
    }


def cmd_create(
    display_name: str,
    fields_json: str,
    unique: str | None = None,
    auto_increment: bool = False,
    table_name: str | None = None,
    description: str = "",
) -> dict:
    # Parse fields
    try:
        fields = json.loads(fields_json)
    except json.JSONDecodeError as e:
        return {"success": False, "error": f"Invalid --fields JSON: {e}"}
    if not isinstance(fields, list) or len(fields) == 0:
        return {"success": False, "error": "--fields must be a non-empty JSON array"}

    # Validate field definitions
    field_names: set[str] = set()
    for i, fdef in enumerate(fields):
        if not isinstance(fdef, dict):
            return {"success": False, "error": f"Field {i}: expected object, got {type(fdef).__name__}"}
        fname = fdef.get("name", "").strip()
        if not fname:
            return {"success": False, "error": f"Field {i}: 'name' is required"}
        if fname in field_names:
            return {"success": False, "error": f"Duplicate field name: '{fname}'"}
        ftype = fdef.get("type", "string")
        if ftype not in VALID_TYPES:
            return {"success": False, "error": f"Field '{fname}': unknown type '{ftype}'. Valid: {', '.join(sorted(VALID_TYPES))}"}
        field_names.add(fname)

    # Parse unique key
    unique_key: list[str] = []
    if unique:
        unique_key = [u.strip() for u in unique.split(",") if u.strip()]
        for uk in unique_key:
            if uk not in field_names and uk != "_id":
                return {"success": False, "error": f"Unique key field '{uk}' not found in field definitions"}

    conn = _get_conn()

    # Check duplicate display_name
    dup = conn.execute("SELECT actual_name FROM _registry WHERE display_name = ?", [display_name]).fetchone()
    if dup:
        return {"success": False, "error": f"A table named '{display_name}' already exists (actual: {dup[0]})."}

    actual_name = table_name if table_name else _generate_table_name(display_name, conn)

    # Build schema fields + DDL columns
    schema_fields: list[dict] = []
    col_defs: list[str] = []

    if auto_increment:
        if "_id" in field_names:
            return {"success": False, "error": "Cannot use --auto-increment: '_id' already in --fields"}
        col_defs.append('"_id" INTEGER')
        schema_fields.append({
            "name": "_id", "type": "integer", "required": False,
            "auto_increment": True, "description": "自动编号",
        })
        if not unique_key:
            unique_key = ["_id"]

    for fdef in fields:
        fname = fdef.get("name", "").strip()
        ftype = fdef.get("type", "string")
        duck_type = TYPE_MAP.get(ftype, "VARCHAR")
        col_defs.append(f'{_q(fname)} {duck_type}')
        schema_fields.append({
            "name": fname, "type": ftype,
            "required": fdef.get("required", False),
            "description": fdef.get("description", ""),
        })

    # Create table
    try:
        conn.execute(f'CREATE TABLE {_q(actual_name)} ({", ".join(col_defs)})')
    except duckdb.Error as e:
        return {"success": False, "error": f"Failed to create table: {e}"}

    if auto_increment:
        conn.execute(f'CREATE SEQUENCE {_q(f"seq_{actual_name}")} START 1')

    # Register
    conn.execute(
        """INSERT INTO _registry
           (actual_name, display_name, description, record_count, created_at, updated_at,
            fields_json, unique_key, auto_increment, auto_increment_field)
           VALUES (?, ?, ?, 0, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP, ?, ?, ?, ?)""",
        [actual_name, display_name, description,
         json.dumps(schema_fields, ensure_ascii=False),
         json.dumps(unique_key), auto_increment,
         "_id" if auto_increment else None],
    )

    return {
        "success": True,
        "message": f"Table '{display_name}' created.",
        "display_name": display_name,
        "actual_name": actual_name,
        "field_count": len(schema_fields),
    }


def cmd_insert(table: str, data_json: str, batch: bool = False) -> dict:
    conn = _get_conn()
    actual = _resolve_table(table, conn)
    if actual is None:
        return {"success": False, "error": f"Table '{table}' not found."}

    reg = conn.execute(
        "SELECT fields_json, unique_key, auto_increment, auto_increment_field FROM _registry WHERE actual_name = ?",
        [actual],
    ).fetchone()
    if reg is None:
        return {"success": False, "error": f"Registry entry missing for '{actual}'."}

    fields = json.loads(reg[0])
    unique_key = json.loads(reg[1])
    auto_inc = reg[2]
    auto_inc_field = reg[3]

    # Parse input
    try:
        raw = json.loads(data_json)
    except json.JSONDecodeError as e:
        return {"success": False, "error": f"Invalid --data JSON: {e}"}
    rows = [raw] if isinstance(raw, dict) else raw
    if not isinstance(rows, list):
        return {"success": False, "error": "--data must be a JSON object or array"}

    # Build a schema dict for _validate_row
    schema = {
        "fields": fields,
        "unique_key": unique_key,
        "auto_increment": auto_inc,
        "auto_increment_field": auto_inc_field,
    }

    inserted: list[dict] = []
    failed: list[dict] = []

    for i, row in enumerate(rows):
        # Assign auto-increment id
        if auto_inc and auto_inc_field and auto_inc_field not in row:
            try:
                next_id = conn.execute(f"SELECT nextval('seq_{actual}')").fetchone()[0]
                row[auto_inc_field] = next_id
            except duckdb.Error as e:
                return {"success": False, "error": f"Failed to get next ID: {e}"}

        errors = _validate_row(row, schema, conn, actual)
        if errors:
            failed.append({"row": i, "errors": errors})
            if not batch:
                return {"success": False, "error": f"Validation failed for row {i}", "details": errors}
            continue

        # Coerce values
        clean_row: dict = {}
        col_names: list[str] = []
        values: list = []
        for fdef in fields:
            fname = fdef["name"]
            ftype = fdef.get("type", "string")
            col_names.append(fname)
            if fname in row:
                val, _ = _coerce_value(row[fname], ftype)
                values.append(val)
            else:
                values.append(None)

        try:
            quoted = [_q(c) for c in col_names]
            placeholders = ", ".join(["?" for _ in col_names])
            conn.execute(
                f'INSERT INTO {_q(actual)} ({", ".join(quoted)}) VALUES ({placeholders})',
                values,
            )
            clean_row = dict(zip(col_names, values))
            inserted.append(clean_row)
        except duckdb.Error as e:
            failed.append({"row": i, "errors": [{"error": str(e)}]})
            if not batch:
                return {"success": False, "error": f"Insert failed for row {i}: {e}"}

    if not inserted:
        return {"success": False, "error": "No valid rows to insert.", "failed": len(failed), "details": failed}

    # Update record count
    count = conn.execute(f'SELECT COUNT(*) FROM {_q(actual)}').fetchone()[0]
    conn.execute(
        "UPDATE _registry SET record_count = ?, updated_at = CURRENT_TIMESTAMP WHERE actual_name = ?",
        [count, actual],
    )

    result: dict = {"success": True, "inserted": len(inserted), "total_rows": count}
    if failed:
        result["failed"] = len(failed)
        result["details"] = failed
    return result


def cmd_update_schema(
    table: str,
    add: str | None = None,
    remove: str | None = None,
    rename: str | None = None,
    modify: str | None = None,
) -> dict:
    conn = _get_conn()
    actual = _resolve_table(table, conn)
    if actual is None:
        return {"success": False, "error": f"Table '{table}' not found."}

    reg = conn.execute(
        "SELECT fields_json, unique_key, auto_increment, auto_increment_field FROM _registry WHERE actual_name = ?",
        [actual],
    ).fetchone()
    if reg is None:
        return {"success": False, "error": f"Registry entry missing for '{actual}'."}

    fields = json.loads(reg[0])
    unique_key = json.loads(reg[1])
    field_map = {f["name"]: f for f in fields}
    changes: list[str] = []

    # ── Add fields ──────────────────────────────────────────────────
    if add:
        try:
            new_fields = json.loads(add)
        except json.JSONDecodeError as e:
            return {"success": False, "error": f"Invalid --add JSON: {e}"}
        if not isinstance(new_fields, list):
            return {"success": False, "error": "--add must be a JSON array"}
        for fdef in new_fields:
            fname = fdef.get("name", "").strip()
            if not fname:
                return {"success": False, "error": "Each added field needs a 'name'"}
            if fname in field_map:
                return {"success": False, "error": f"Field '{fname}' already exists"}
            ftype = fdef.get("type", "string")
            if ftype not in VALID_TYPES:
                return {"success": False, "error": f"Field '{fname}': unknown type '{ftype}'"}
            duck_type = TYPE_MAP.get(ftype, "VARCHAR")
            try:
                conn.execute(f'ALTER TABLE {_q(actual)} ADD COLUMN {_q(fname)} {duck_type}')
            except duckdb.Error as e:
                return {"success": False, "error": f"Failed to add column: {e}"}
            new_fdef = {"name": fname, "type": ftype, "required": fdef.get("required", False),
                        "description": fdef.get("description", "")}
            fields.append(new_fdef)
            field_map[fname] = new_fdef
            changes.append(f"added field '{fname}' ({ftype})")

    # ── Remove fields ───────────────────────────────────────────────
    if remove:
        names = [n.strip() for n in remove.split(",") if n.strip()]
        for fname in names:
            if fname not in field_map:
                return {"success": False, "error": f"Field '{fname}' does not exist"}
            if fname == "_id" and reg[2]:
                return {"success": False, "error": "Cannot remove auto-increment field '_id'"}
            if fname in unique_key:
                return {"success": False, "error": f"Cannot remove unique key field '{fname}'"}
            try:
                conn.execute(f'ALTER TABLE {_q(actual)} DROP COLUMN {_q(fname)}')
            except duckdb.Error as e:
                return {"success": False, "error": f"Failed to drop column: {e}"}
            fields = [f for f in fields if f["name"] != fname]
            del field_map[fname]
            changes.append(f"removed field '{fname}'")

    # ── Rename field ────────────────────────────────────────────────
    if rename:
        parts = rename.split(":", 1)
        if len(parts) != 2:
            return {"success": False, "error": "--rename format: old_name:new_name"}
        old_name, new_name = parts[0].strip(), parts[1].strip()
        if old_name not in field_map:
            return {"success": False, "error": f"Field '{old_name}' does not exist"}
        if new_name in field_map:
            return {"success": False, "error": f"Field '{new_name}' already exists"}
        try:
            conn.execute(f'ALTER TABLE {_q(actual)} RENAME COLUMN {_q(old_name)} TO {_q(new_name)}')
        except duckdb.Error as e:
            return {"success": False, "error": f"Failed to rename column: {e}"}
        fdef = field_map.pop(old_name)
        fdef["name"] = new_name
        field_map[new_name] = fdef
        if old_name in unique_key:
            unique_key = [new_name if k == old_name else k for k in unique_key]
        # Update field list
        for f in fields:
            if f["name"] == old_name:
                f["name"] = new_name
        changes.append(f"renamed field '{old_name}' → '{new_name}'")

    # ── Modify field type ───────────────────────────────────────────
    if modify:
        try:
            mod_fields = json.loads(modify)
        except json.JSONDecodeError as e:
            return {"success": False, "error": f"Invalid --modify JSON: {e}"}
        if not isinstance(mod_fields, list):
            return {"success": False, "error": "--modify must be a JSON array"}
        for mdef in mod_fields:
            mname = mdef.get("name", "").strip()
            if mname not in field_map:
                return {"success": False, "error": f"Field '{mname}' does not exist"}
            mtype = mdef.get("type", "string")
            if mtype not in VALID_TYPES:
                return {"success": False, "error": f"Field '{mname}': unknown type '{mtype}'"}
            new_type = TYPE_MAP.get(mtype, "VARCHAR")
            old_type = field_map[mname]["type"]
            try:
                conn.execute(f'ALTER TABLE {_q(actual)} ALTER COLUMN {_q(mname)} TYPE {new_type}')
            except duckdb.Error as e:
                return {"success": False, "error": f"Failed to alter column type: {e}"}
            field_map[mname]["type"] = mtype
            for f in fields:
                if f["name"] == mname:
                    f["type"] = mtype
            changes.append(f"modified field '{mname}': {old_type} → {mtype}")

    if not changes:
        return {"success": False, "error": "No changes specified."}

    # Save updated fields_json and unique_key
    conn.execute(
        "UPDATE _registry SET fields_json = ?, unique_key = ?, updated_at = CURRENT_TIMESTAMP WHERE actual_name = ?",
        [json.dumps(fields, ensure_ascii=False), json.dumps(unique_key), actual],
    )

    # Update record count
    try:
        count = conn.execute(f'SELECT COUNT(*) FROM {_q(actual)}').fetchone()[0]
        conn.execute("UPDATE _registry SET record_count = ? WHERE actual_name = ?", [count, actual])
    except duckdb.Error:
        pass

    return {"success": True, "message": "Schema updated.", "changes": changes}


def cmd_delete(table: str, keep_files: bool = False) -> dict:
    conn = _get_conn()
    actual = _resolve_table(table, conn)
    if actual is None:
        return {"success": False, "error": f"Table '{table}' not found."}

    display_name = conn.execute(
        "SELECT display_name FROM _registry WHERE actual_name = ?", [actual]
    ).fetchone()[0]

    # Drop table and sequence
    if not keep_files:
        try:
            conn.execute(f'DROP TABLE IF EXISTS {_q(actual)}')
        except duckdb.Error:
            pass
        try:
            conn.execute(f'DROP SEQUENCE IF EXISTS {_q(f"seq_{actual}")}')
        except duckdb.Error:
            pass
        conn.execute("DELETE FROM _embeddings WHERE table_name = ?", [actual])

    conn.execute("DELETE FROM _registry WHERE actual_name = ?", [actual])
    return {"success": True, "message": f"Table '{display_name}' deleted.", "actual_name": actual}


def cmd_stats(table: str | None = None) -> dict:
    conn = _get_conn()
    if table:
        actual = _resolve_table(table, conn)
        if actual is None:
            return {"success": False, "error": f"Table '{table}' not found."}
        reg = conn.execute(
            "SELECT display_name, fields_json, record_count, created_at, updated_at "
            "FROM _registry WHERE actual_name = ?", [actual]
        ).fetchone()
        if reg is None:
            return {"success": False, "error": f"Registry entry missing for '{actual}'."}
        fields = json.loads(reg[1])
        return {
            "success": True,
            "display_name": reg[0],
            "actual_name": actual,
            "field_count": len(fields),
            "row_count": reg[2],
            "created_at": reg[3].isoformat() if reg[3] else "",
            "updated_at": reg[4].isoformat() if reg[4] else "",
        }

    # All tables
    rows = conn.execute(
        "SELECT display_name, actual_name, record_count, updated_at FROM _registry ORDER BY display_name"
    ).fetchall()
    tables = []
    total_rows = 0
    for r in rows:
        tables.append({
            "display_name": r[0],
            "actual_name": r[1],
            "row_count": r[2],
            "updated_at": r[3].isoformat() if r[3] else "",
        })
        total_rows += r[2]
    return {"success": True, "table_count": len(tables), "total_rows": total_rows, "tables": tables}


def cmd_embed(table: str | None = None) -> dict:
    """Generate embeddings for table rows and store in _embeddings table."""
    conn = _get_conn()

    try:
        from generate_embeddings import get_embedding
    except ImportError:
        return {"success": False, "error": "Embedding module not available. Install sentence-transformers."}

    if table:
        actual = _resolve_table(table, conn)
        if actual is None:
            return {"success": False, "error": f"Table '{table}' not found."}
        table_names = [actual]
    else:
        table_names = [r[0] for r in conn.execute("SELECT actual_name FROM _registry").fetchall()]

    total_embedded = 0
    for actual_name in table_names:
        reg = conn.execute(
            "SELECT fields_json, display_name FROM _registry WHERE actual_name = ?", [actual_name]
        ).fetchone()
        if reg is None:
            continue
        fields = json.loads(reg[0])
        text_cols = [f["name"] for f in fields if f.get("type") in ("string", "text") and not f.get("auto_increment")]
        if not text_cols:
            continue

        try:
            rows = conn.execute(f'SELECT * FROM {_q(actual_name)}').fetchall()
            col_names = [desc[0] for desc in conn.description]
        except duckdb.Error:
            continue

        embedded = 0
        for row in rows:
            row_dict = dict(zip(col_names, row))
            row_id = row_dict.get("_id", 0)
            text_parts = [str(row_dict.get(c, "")) for c in text_cols if row_dict.get(c) is not None]
            full_text = " ".join(text_parts)
            if not full_text.strip():
                continue

            emb = get_embedding(full_text)
            if emb is None:
                continue

            # Store as DuckDB array
            emb_str = "[" + ", ".join(str(v) for v in emb) + "]"
            try:
                conn.execute(
                    f"""INSERT INTO _embeddings (table_name, row_id, embedding, updated_at)
                        VALUES (?, ?, {emb_str}::FLOAT[{len(emb)}], CURRENT_TIMESTAMP)
                        ON CONFLICT (table_name, row_id) DO UPDATE SET
                            embedding = EXCLUDED.embedding, updated_at = CURRENT_TIMESTAMP""",
                    [actual_name, row_id],
                )
                embedded += 1
            except duckdb.Error:
                pass

        total_embedded += embedded

    return {"success": True, "message": f"Embedded {total_embedded} rows.", "embedded": total_embedded}


def _read_tabular_file(filepath: str) -> list[list[str]]:
    path = Path(filepath)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {filepath}")

    ext = path.suffix.lower()
    if ext == ".csv":
        with open(path, encoding="utf-8-sig", errors="replace", newline="") as f:
            return list(csv.reader(f))

    if ext in (".xlsx", ".xls"):
        try:
            import openpyxl
        except ImportError as exc:
            raise RuntimeError("openpyxl not installed. Run: pip install openpyxl") from exc
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.active
            return [["" if c.value is None else str(c.value) for c in row] for row in ws.iter_rows()]
        finally:
            wb.close()

    with open(path, encoding="utf-8-sig", errors="replace", newline="") as f:
        return list(csv.reader(f))


def _next_display_name(display_name: str, conn: duckdb.DuckDBPyConnection) -> str:
    row = conn.execute("SELECT 1 FROM _registry WHERE display_name = ?", [display_name]).fetchone()
    if not row:
        return display_name
    suffix = 2
    while True:
        candidate = f"{display_name} ({suffix})"
        row = conn.execute("SELECT 1 FROM _registry WHERE display_name = ?", [candidate]).fetchone()
        if not row:
            return candidate
        suffix += 1


def cmd_import(filepath: str, table_name: Optional[str] = None) -> dict:
    """Import CSV/Excel data into the DuckDB ledger backend."""
    raw_rows = _read_tabular_file(filepath)
    if not raw_rows or len(raw_rows) < 2:
        return {"success": False, "error": "Table must have at least 1 header row and 1 data row"}

    headers = _dedupe_headers(raw_rows[0])
    if not any(h.strip() for h in headers):
        return {"success": False, "error": "No headers found in first row"}

    data_rows = raw_rows[1:]
    sample_rows = data_rows[:100]
    fields = []
    for i, header in enumerate(headers):
        values = [row[i] for row in sample_rows if i < len(row)]
        fields.append({
            "name": header,
            "type": _infer_import_type(values),
            "required": False,
            "description": "",
        })

    source = Path(filepath).name
    display_name = table_name or source.replace("_", " ").replace("-", " ").rsplit(".", 1)[0]

    conn = _get_conn()
    display_name = _next_display_name(display_name, conn)
    actual_name = _generate_table_name(display_name, conn)
    create_result = cmd_create(
        display_name=display_name,
        fields_json=json.dumps(fields, ensure_ascii=False),
        table_name=actual_name,
        description=f"Table imported from {source}",
    )
    if not create_result.get("success"):
        return create_result

    rows = []
    for raw_row in data_rows:
        row_data = {}
        for i, field in enumerate(fields):
            raw_value = raw_row[i] if i < len(raw_row) else ""
            row_data[field["name"]] = _clean_import_value(raw_value, field["type"])
        rows.append(row_data)

    insert_result = cmd_insert(actual_name, json.dumps(rows, ensure_ascii=False), batch=True)
    if not insert_result.get("success"):
        cmd_delete(actual_name)
        return {"success": False, "error": "Import failed during insert", "details": insert_result}

    result = {
        "success": True,
        "message": f"Imported '{source}' as '{display_name}'.",
        "display_name": display_name,
        "actual_name": actual_name,
        "source": source,
        "row_count": insert_result.get("inserted", 0),
        "field_count": len(fields),
        "fields": fields,
    }
    if insert_result.get("failed"):
        result["failed"] = insert_result["failed"]
        result["details"] = insert_result.get("details", [])
    return result


def cmd_export(table: str, output_path: Optional[str] = None) -> dict:
    """Export a DuckDB ledger table as CSV."""
    try:
        conn = _get_conn(read_only=True)
    except FileNotFoundError:
        return {"success": False, "error": f"Table '{table}' not found."}
    actual = _resolve_table(table, conn)
    if actual is None:
        return {"success": False, "error": f"Table '{table}' not found."}

    reg = conn.execute(
        "SELECT display_name, fields_json FROM _registry WHERE actual_name = ?", [actual]
    ).fetchone()
    if reg is None:
        return {"success": False, "error": f"Registry entry missing for '{actual}'."}

    fields = json.loads(reg[1])
    col_names = [f["name"] for f in fields]
    output = output_path or f"{actual}.csv"
    rows = conn.execute(f'SELECT {", ".join(_q(c) for c in col_names)} FROM {_q(actual)}').fetchall()

    with open(output, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(col_names)
        writer.writerows(rows)

    return {
        "success": True,
        "message": f"Exported '{reg[0]}' to {output}.",
        "output": output,
        "row_count": len(rows),
    }


def search_ledgers(query: str, limit: int = 10) -> list[dict]:
    """Search across DuckDB ledger tables by name, field name, or cell content."""
    try:
        conn = _get_conn(read_only=True)
    except FileNotFoundError:
        return []
    q = query.lower()
    results: list[dict] = []

    registry_rows = conn.execute(
        "SELECT actual_name, display_name, description, record_count, fields_json, updated_at "
        "FROM _registry"
    ).fetchall()
    for actual, display, description, record_count, fields_json, updated_at in registry_rows:
        fields = json.loads(fields_json)
        score = 0
        matched_fields = []
        if q in (display or "").lower():
            score += 10
        if q in (description or "").lower():
            score += 4
        for field in fields:
            fname = field["name"]
            if q in fname.lower():
                score += 3
                matched_fields.append(f"{fname} ({field.get('type', 'text')})")

        preview = []
        try:
            rows = conn.execute(f'SELECT * FROM {_q(actual)} LIMIT 50').fetchall()
            col_names = [d[0] for d in conn.description]
        except duckdb.Error:
            rows = []
            col_names = []

        content_matches = []
        for row in rows:
            row_dict = dict(zip(col_names, row))
            if any(q in str(v).lower() for v in row_dict.values() if v is not None):
                score += 1
                content_matches.append(row_dict)
        if content_matches:
            preview = content_matches[:3]
        elif score > 0:
            preview = [dict(zip(col_names, row)) for row in rows[:3]]

        if score > 0:
            results.append({
                "id": actual,
                "name": display,
                "type": "ledger",
                "score": score,
                "fields": [f["name"] for f in fields],
                "field_types": {f["name"]: f.get("type", "text") for f in fields},
                "rows_count": record_count,
                "matched_fields": matched_fields,
                "preview": preview,
                "updated_at": updated_at.isoformat() if updated_at else "",
            })

    results.sort(key=lambda item: -item["score"])
    return results[:limit]


def cmd_search(query: str, limit: int = 10) -> dict:
    results = search_ledgers(query, limit=limit)
    return {"success": True, "count": len(results), "results": results}


# ═══════════════════════════════════════════════════════════════════════
# CLI
# ═══════════════════════════════════════════════════════════════════════

def main():
    parser = argparse.ArgumentParser(description="Ledger/台账 management for LLM Wiki v2 (DuckDB backend)")
    subparsers = parser.add_subparsers(dest="command", required=True)

    subparsers.add_parser("list", help="List all tables")

    show_parser = subparsers.add_parser("show", help="Show table schema and data")
    show_parser.add_argument("table", help="Table name (display or actual)")

    create_parser = subparsers.add_parser("create", help="Create a new table")
    create_parser.add_argument("display_name", help="Display name for the table")
    create_parser.add_argument("--fields", required=True, help='Field definitions JSON')
    create_parser.add_argument("--unique", default=None, help="Unique key field(s)")
    create_parser.add_argument("--auto-increment", action="store_true", help="Add _id auto-increment field")
    create_parser.add_argument("--table-name", default=None, help="Override safe table name")
    create_parser.add_argument("--description", default="", help="Table description")

    insert_parser = subparsers.add_parser("insert", help="Insert data into a table")
    insert_parser.add_argument("table", help="Table name")
    insert_parser.add_argument("--data", required=True, help="JSON data (object or array)")
    insert_parser.add_argument("--batch", action="store_true", help="Continue on partial errors")

    update_parser = subparsers.add_parser("update-schema", help="Modify table schema")
    update_parser.add_argument("table", help="Table name")
    update_parser.add_argument("--add", default=None, help="Add fields JSON")
    update_parser.add_argument("--remove", default=None, help="Remove fields: name1,name2")
    update_parser.add_argument("--rename", default=None, help="Rename field: old:new")
    update_parser.add_argument("--modify", default=None, help="Change field type JSON")

    del_parser = subparsers.add_parser("delete", help="Delete a table")
    del_parser.add_argument("table", help="Table name")
    del_parser.add_argument("--keep-files", action="store_true", help="Keep data in DuckDB (remove only registry entry)")

    stats_parser = subparsers.add_parser("stats", help="Show table statistics")
    stats_parser.add_argument("table", nargs="?", default=None, help="Table name (omit for all)")

    embed_parser = subparsers.add_parser("embed", help="Generate embeddings for table rows")
    embed_parser.add_argument("table", nargs="?", default=None, help="Table name (omit for all)")

    import_parser = subparsers.add_parser("import", help="Import CSV/Excel as a ledger table")
    import_parser.add_argument("file", help="CSV or XLSX file path")
    import_parser.add_argument("--name", help="Table display name")

    search_parser = subparsers.add_parser("search", help="Search across ledger tables")
    search_parser.add_argument("query", help="Search query")
    search_parser.add_argument("-n", "--limit", type=int, default=10, help="Max results")

    export_parser = subparsers.add_parser("export", help="Export a ledger table as CSV")
    export_parser.add_argument("table", help="Table name (display or actual)")
    export_parser.add_argument("-o", "--output", help="Output CSV path")

    args = parser.parse_args()

    if args.command == "list":
        result = cmd_list()
    elif args.command == "show":
        result = cmd_show(args.table)
    elif args.command == "create":
        result = cmd_create(
            display_name=args.display_name, fields_json=args.fields, unique=args.unique,
            auto_increment=args.auto_increment, table_name=args.table_name,
            description=args.description,
        )
    elif args.command == "insert":
        result = cmd_insert(args.table, args.data, batch=args.batch)
    elif args.command == "update-schema":
        result = cmd_update_schema(args.table, add=args.add, remove=args.remove,
                                   rename=args.rename, modify=args.modify)
    elif args.command == "delete":
        result = cmd_delete(args.table, keep_files=args.keep_files)
    elif args.command == "stats":
        result = cmd_stats(args.table)
    elif args.command == "embed":
        result = cmd_embed(args.table)
    elif args.command == "import":
        result = cmd_import(args.file, table_name=args.name)
    elif args.command == "search":
        result = cmd_search(args.query, limit=args.limit)
    elif args.command == "export":
        result = cmd_export(args.table, output_path=args.output)
    else:
        result = {"success": False, "error": f"Unknown command: {args.command}"}

    print(json.dumps(result, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
